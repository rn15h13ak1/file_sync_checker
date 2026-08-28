"""reporter: Excel/HTML 出力スモークテスト。

I/O ロジックなので「正しく生成され、想定のシート/セクション・色塗り・プレースホルダが入る」
ことを確認する。色やレイアウトの細部までは確認しない。
"""
from __future__ import annotations

import html as html_mod
import json
import re
import tracemalloc
import unicodedata
from datetime import datetime
from pathlib import Path

import pytest

from comparator import (
    STATUS_ERROR,
    STATUS_HASH_MISMATCH,
    STATUS_OK,
    STATUS_PARTIAL_MISSING,
    STATUS_PARTIAL_PRESENT,
    compare,
)
from reporter import (
    ERROR_PLACEHOLDER,
    MISSING_PLACEHOLDER,
    SKIPPED_HASH_PLACEHOLDER,
    ReportContext,
    _build_paths,
    _looks_windows,
    write_excel,
    write_html,
)

from .conftest import make_entry, make_scan


# ============================================================
# パス組み立てヘルパー
# ============================================================
class TestLooksWindows:
    def test_unc_forward_slash(self):
        assert _looks_windows("//server/share/docs") is True

    def test_unc_backslash(self):
        assert _looks_windows("\\\\server\\share\\docs") is True

    def test_drive_letter(self):
        assert _looks_windows("C:/Users/foo") is True
        assert _looks_windows("D:") is True

    def test_unix_path(self):
        assert _looks_windows("/mnt/share/docs") is False
        assert _looks_windows("/Users/foo") is False

    def test_relative_path(self):
        assert _looks_windows("relative/path") is False


class TestBuildPaths:
    def test_unc_forward_slash_normalizes_to_backslash(self):
        """設定が //server/share/docs でも Windows 形式に統一する。"""
        full, folder = _build_paths("//server-a/share/docs", "設計/詳細/spec.xlsx")
        assert full == "\\\\server-a\\share\\docs\\設計\\詳細\\spec.xlsx"
        assert folder == "\\\\server-a\\share\\docs\\設計\\詳細"

    def test_file_at_root_of_share(self):
        full, folder = _build_paths("//server/share", "file.txt")
        assert full == "\\\\server\\share\\file.txt"
        assert folder == "\\\\server\\share"

    def test_unix_root_keeps_forward_slash(self):
        full, folder = _build_paths("/mnt/data", "sub/file.txt")
        assert full == "/mnt/data/sub/file.txt"
        assert folder == "/mnt/data/sub"

    def test_drive_letter_root(self):
        full, folder = _build_paths("C:/Users/foo", "docs/file.txt")
        assert full == "C:\\Users\\foo\\docs\\file.txt"
        assert folder == "C:\\Users\\foo\\docs"

    def test_trailing_separator_stripped(self):
        """ルートの末尾セパレータがあっても二重にならない。"""
        full, _ = _build_paths("//server/share/", "file.txt")
        assert "\\\\\\" not in full  # 三重セパレータが現れない
        assert full == "\\\\server\\share\\file.txt"


# ============================================================
# 共通: 全状態を含むコンテキストを作る
# ============================================================
@pytest.fixture
def rich_ctx(tmp_path: Path) -> ReportContext:
    """OK / 不一致 / 欠落 / 一部のみ / エラー / dir差分 を1件ずつ含むコンテキスト。"""
    # 3拠点
    a = make_scan(
        "拠点A",
        files={
            "ok.txt": make_entry("hash_ok"),
            "mismatch.txt": make_entry("hash_majority"),    # A=B≠C
            "missing_in_C.txt": make_entry("hash_x"),
            "only_in_A.txt": make_entry("hash_y"),
        },
        dirs={"only_A_dir": object()},
    )
    b = make_scan(
        "拠点B",
        files={
            "ok.txt": make_entry("hash_ok"),
            "mismatch.txt": make_entry("hash_majority"),
            "missing_in_C.txt": make_entry("hash_x"),
        },
    )
    c = make_scan(
        "拠点C",
        files={
            "ok.txt": make_entry("hash_ok"),
            "mismatch.txt": make_entry("hash_minority"),    # ← 少数派
        },
        file_errors={"locked.txt": "PermissionError: [Errno 13]"},
    )
    comparison = compare([a, b, c])

    return ReportContext(
        started_at=datetime(2026, 5, 20, 12, 0, 0),
        finished_at=datetime(2026, 5, 20, 12, 0, 3),
        config_path=tmp_path / "fake.yaml",
        scans=[a, b, c],
        comparison=comparison,
    )


# ============================================================
# Excel
# ============================================================
class TestExcel:
    def test_writes_xlsx_with_expected_sheets(self, rich_ctx: ReportContext, tmp_path: Path):
        out = write_excel(rich_ctx, tmp_path / "out.xlsx")
        assert out.is_file()
        assert out.stat().st_size > 0

        from openpyxl import load_workbook
        wb = load_workbook(out)
        expected = {
            "サマリー",
            "全ファイル一覧",
            "ハッシュ不一致",
            "ファイル欠落",
            "余分なファイル",
            "エラー対象ファイル",
            "フォルダ構造差分",
            "エラー",
        }
        assert expected.issubset(set(wb.sheetnames)), wb.sheetnames

    def test_all_files_sheet_has_one_row_per_relpath(self, rich_ctx: ReportContext, tmp_path: Path):
        out = write_excel(rich_ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        ws = wb["全ファイル一覧"]
        # ヘッダ + 5ファイル (ok, mismatch, missing_in_C, only_in_A, locked)
        assert ws.max_row == 1 + 5

        # 列: No, relpath, name, ext, status, [A:size,A:hash,A:mtime, B:..., C:...]
        # = 5 + 3*3 = 14
        assert ws.max_column == 14

    def test_status_column_values(self, rich_ctx: ReportContext, tmp_path: Path):
        out = write_excel(rich_ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        ws = wb["全ファイル一覧"]
        statuses = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=5).value
                    for r in range(2, ws.max_row + 1)}
        assert statuses["ok.txt"] == STATUS_OK
        assert statuses["mismatch.txt"] == STATUS_HASH_MISMATCH
        assert statuses["missing_in_C.txt"] == STATUS_PARTIAL_MISSING
        assert statuses["only_in_A.txt"] == STATUS_PARTIAL_PRESENT
        assert statuses["locked.txt"] == STATUS_ERROR

    def test_minority_hash_highlighting(self, rich_ctx: ReportContext, tmp_path: Path):
        """A=B≠C のとき C のハッシュ列だけ赤背景になる (Bug A 修正)。"""
        out = write_excel(rich_ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        ws = wb["全ファイル一覧"]
        # mismatch.txt の行を探す
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=2).value == "mismatch.txt":
                # 拠点A:hash=列7, 拠点B:hash=列10, 拠点C:hash=列13
                a_fill = ws.cell(row=r, column=7).fill.fgColor.rgb
                b_fill = ws.cell(row=r, column=10).fill.fgColor.rgb
                c_fill = ws.cell(row=r, column=13).fill.fgColor.rgb
                # 多数派 (A, B) は塗りなし
                assert a_fill in ("00000000", None)
                assert b_fill in ("00000000", None)
                # 少数派 (C) は赤 FILL_HASH_MISMATCH = FFC7CE
                assert c_fill == "00FFC7CE", f"expected red, got {c_fill}"
                return
        pytest.fail("mismatch.txt row not found")

    def test_missing_cells_use_half_width_placeholder(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """欠落セルは半角 '-' (項目 P 修正)。"""
        out = write_excel(rich_ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        ws = wb["全ファイル一覧"]
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=2).value == "only_in_A.txt":
                # 拠点B, C は欠落 → 列 9,10,11 と 12,13,14
                for c in (9, 10, 11, 12, 13, 14):
                    assert ws.cell(row=r, column=c).value == MISSING_PLACEHOLDER
                assert MISSING_PLACEHOLDER == "-"
                return
        pytest.fail("only_in_A.txt row not found")

    def test_error_row_renders_error_placeholder_and_message(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """エラー対象ファイル: サイズ列に 'エラー'、ハッシュ列にメッセージ。"""
        out = write_excel(rich_ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        ws = wb["全ファイル一覧"]
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=2).value == "locked.txt":
                # 拠点C: 列12=size, 13=hash, 14=mtime
                assert ws.cell(row=r, column=12).value == ERROR_PLACEHOLDER
                assert "PermissionError" in str(ws.cell(row=r, column=13).value)
                return
        pytest.fail("locked.txt row not found")

    def test_subset_sheets_only_contain_matching_rows(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        out = write_excel(rich_ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        # ハッシュ不一致シートには mismatch.txt のみ
        mismatch_sheet = wb["ハッシュ不一致"]
        rels = [mismatch_sheet.cell(row=r, column=2).value
                for r in range(2, mismatch_sheet.max_row + 1)]
        assert rels == ["mismatch.txt"]

        # ファイル欠落: missing_in_C.txt のみ
        missing_sheet = wb["ファイル欠落"]
        rels = [missing_sheet.cell(row=r, column=2).value
                for r in range(2, missing_sheet.max_row + 1)]
        assert rels == ["missing_in_C.txt"]

        # 余分なファイル: only_in_A.txt のみ
        extra_sheet = wb["余分なファイル"]
        rels = [extra_sheet.cell(row=r, column=2).value
                for r in range(2, extra_sheet.max_row + 1)]
        assert rels == ["only_in_A.txt"]

        # エラー対象ファイル: locked.txt のみ
        errored_sheet = wb["エラー対象ファイル"]
        rels = [errored_sheet.cell(row=r, column=2).value
                for r in range(2, errored_sheet.max_row + 1)]
        assert rels == ["locked.txt"]

    def test_empty_subset_shows_placeholder(self, tmp_path: Path):
        """差分が無いケースでは '該当なし' が入る。"""
        a = make_scan("A", files={"x.txt": make_entry("h")})
        b = make_scan("B", files={"x.txt": make_entry("h")})
        comparison = compare([a, b])
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=comparison,
        )
        out = write_excel(ctx, tmp_path / "empty.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        for sheet in ("ハッシュ不一致", "ファイル欠落", "余分なファイル", "エラー対象ファイル"):
            ws = wb[sheet]
            assert ws.cell(row=2, column=1).value == "該当なし", sheet

    def test_panes_and_autofilter_survive_write_only(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """見出し行の固定とオートフィルタが保存後も残る。

        write_only ではシートビューが行より先に出力されるため、freeze_panes を
        append より後に設定すると黙って捨てられる。
        """
        out = write_excel(rich_ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)

        for sheet in ("全ファイル一覧", "ハッシュ不一致", "フォルダ構造差分"):
            assert wb[sheet].freeze_panes == "C2", sheet
        assert wb["エラー"].freeze_panes == "A2"

        ws = wb["全ファイル一覧"]
        # ヘッダー + 5ファイル、3拠点 = 14列
        assert ws.auto_filter.ref == "A1:N6"

    def test_truncates_at_excel_row_limit(self, tmp_path: Path, monkeypatch):
        """行数上限を超えたら切り詰めて注記する。

        上限超過で保存が失敗すると Excel レポートが 1 枚も残らないため、
        全滅させるより切り詰めて HTML 側へ誘導する。
        実際の上限は 100 万行なので、テストでは上限を差し替えて確認する。
        """
        import reporter

        monkeypatch.setattr(reporter, "EXCEL_MAX_DATA_ROWS", 3)
        files = {f"f{i}.txt": make_entry("h") for i in range(5)}
        a = make_scan("A", files=files)
        b = make_scan("B", files=files)
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=compare([a, b]),
        )
        out = write_excel(ctx, tmp_path / "big.xlsx")

        from openpyxl import load_workbook
        ws = load_workbook(out)["全ファイル一覧"]
        # ヘッダー + データ3行 + 注記1行
        assert ws.max_row == 5
        assert "他 2 件" in str(ws.cell(row=5, column=1).value)

    def test_error_sheet_lists_scan_errors(self, tmp_path: Path):
        """エラーシートに拠点・相対パス・メッセージが並ぶ。"""
        from scanner import ScanError

        a = make_scan(
            "拠点A",
            files={"x.txt": make_entry("h")},
            errors=[
                ScanError(relpath="", message="root path does not exist: //srv/share"),
                ScanError(relpath="部署/資料.xlsx", message="PermissionError: [Errno 13]"),
            ],
        )
        b = make_scan(
            "拠点B",
            files={"x.txt": make_entry("h")},
            errors=[ScanError(relpath="tmp", message="walk error: [Errno 13]")],
        )
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=compare([a, b]),
        )
        out = write_excel(ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        ws = load_workbook(out)["エラー"]

        assert ws.max_row == 1 + 3  # ヘッダ + 3件
        rows = [
            tuple(ws.cell(row=r, column=c).value for c in range(1, 5))
            for r in range(2, ws.max_row + 1)
        ]
        # ルート自体のエラーは relpath が空。Excel では空文字が空セルになる
        assert rows[0] == (1, "拠点A", None, "root path does not exist: //srv/share")
        assert rows[1] == (2, "拠点A", "部署/資料.xlsx", "PermissionError: [Errno 13]")
        assert rows[2] == (3, "拠点B", "tmp", "walk error: [Errno 13]")
        # 見出し行の固定はエラーシートだけ A2 (相対パス列を固定しない)
        assert ws.freeze_panes == "A2"

    def test_subset_sheet_truncates_at_row_limit(self, tmp_path: Path, monkeypatch):
        """差分シート側でも行数上限で切り詰めて注記する。"""
        import reporter

        monkeypatch.setattr(reporter, "EXCEL_MAX_DATA_ROWS", 2)
        a = make_scan("A", files={f"f{i}.txt": make_entry(f"h{i}") for i in range(5)})
        b = make_scan("B", files={f"f{i}.txt": make_entry(f"other{i}") for i in range(5)})
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=compare([a, b]),
        )
        out = write_excel(ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        ws = load_workbook(out)["ハッシュ不一致"]

        # ヘッダ + データ2行 + 注記1行
        assert ws.max_row == 4
        assert "他 3 件" in str(ws.cell(row=4, column=1).value)

    def test_summary_sheet_includes_diff_counts(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        out = write_excel(rich_ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        ws = wb["サマリー"]
        # 全セル走査でラベル+件数を回収
        all_cells = {}
        for row in ws.iter_rows(values_only=True):
            if len(row) >= 2 and row[0] is not None:
                all_cells[row[0]] = row[1]
        assert all_cells.get("ハッシュ不一致") == 1
        assert all_cells.get("ファイル欠落 (一部拠点になし)") == 1
        assert all_cells.get("余分なファイル (一部拠点のみ)") == 1
        assert all_cells.get("エラー対象ファイル (要確認)") == 1


# ============================================================
# HTML
# ============================================================
def _big_ctx(tmp_path: Path, n_rows: int) -> ReportContext:
    """メモリ計測用に、そこそこ大きいレポートのコンテキストを作る。"""
    files = {
        f"dir{i // 50}/file_{i:05d}_資料.docx": make_entry(f"{i:064x}", size=1234)
        for i in range(n_rows)
    }
    scans = [make_scan(name, files=dict(files)) for name in ("拠点A", "拠点B", "拠点C")]
    return ReportContext(
        started_at=datetime(2026, 1, 1),
        finished_at=datetime(2026, 1, 1),
        config_path=tmp_path / "c.yaml",
        scans=scans,
        comparison=compare(scans),
    )


def _extract_report_data(out: Path) -> dict:
    """レポート HTML に埋め込まれた #report-data の JSON を取り出す。"""
    body = out.read_text(encoding="utf-8")
    m = re.search(
        r'<script type="application/json" id="report-data">\n(.*?)\n</script>',
        body,
        re.S,
    )
    assert m, "#report-data ブロックが見つからない"
    return json.loads(m.group(1))


class TestHtml:
    def test_writes_valid_html(self, rich_ctx: ReportContext, tmp_path: Path):
        out = write_html(rich_ctx, tmp_path / "out.html")
        assert out.is_file()
        body = out.read_text(encoding="utf-8")
        assert body.startswith("\n<!DOCTYPE html>") or body.lstrip().startswith("<!DOCTYPE html>")
        assert "</html>" in body

    def test_contains_all_sections(self, rich_ctx: ReportContext, tmp_path: Path):
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        for sid in ("summary", "all", "mismatch", "missing", "extra", "errored", "dirs", "errors"):
            assert f'id="{sid}"' in body, f"missing section: {sid}"

    def test_status_classes_applied(self, rich_ctx: ReportContext, tmp_path: Path):
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        for cls in (
            "status-ok",
            "status-mismatch",
            "status-missing",
            "status-partial",
            "status-error",
        ):
            assert cls in body, f"missing class: {cls}"

    def test_minority_hash_uses_cell_mismatch_class(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """A=B≠C のとき少数派ハッシュだけが cell-mismatch クラスを持つはず。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        # ハッシュ値は最初の12文字で短縮表示される
        # hash_majority → "hash_majorit…", hash_minority → "hash_minorit…"
        # cell-mismatch クラスが付くのは hash_minority 側のみ
        # 簡易チェック: cell-mismatch がページに存在
        assert "cell-mismatch" in body

    def test_missing_uses_half_width_placeholder(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        # 全角 '－' は使わない
        assert "－" not in body
        # 半角 '-' の欠落セルが存在
        assert "cell-missing'>-" in body

    def test_error_cell_shows_message_in_tooltip(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        # エラーメッセージが title 属性に含まれる
        assert "PermissionError" in body
        assert "cell-error" in body

    def test_sticky_column_has_explicit_background(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """Issue 1: 横スクロール時に sticky 列が透明になる問題が再発しないこと。

        - background: inherit を使っていない
        - thead 側に明示的な背景色
        - tbody の odd/even 両方に明示的な背景色
        """
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        assert "background: inherit" not in body, (
            "sticky 列に background: inherit を使うと透明になる。明示的な色を指定すること"
        )
        # ファイル名列 (3列目) を sticky 対象とする
        assert ".fixed-col-table thead th:nth-child(3)" in body
        # tbody 側の odd/even
        assert ".fixed-col-table tbody tr:nth-child(odd) td:nth-child(3)" in body
        assert ".fixed-col-table tbody tr:nth-child(even) td:nth-child(3)" in body
        # 旧仕様 (相対パス列 = 2 列目) が残っていないこと
        assert ".fixed-col-table tbody td:nth-child(2)" not in body

    def test_page_header_is_not_sticky(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """Issue 2: ページヘッダーを sticky にするとテーブル列ヘッダーと衝突するため、
        ページヘッダー側を解除した状態を維持する。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        # 旧バグの組み合わせを直接禁止
        assert "header { background: #305496; color: #fff; padding: 16px 24px; position: sticky" not in body

    def test_section_has_scroll_margin_top(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """Issue 3: アンカージャンプ時、section 見出しが画面端に潜らないように余白がある。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        assert "scroll-margin-top" in body, "section に scroll-margin-top を設定すること"

    def test_file_row_has_required_data_attributes(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """行クリックでモーダルを開くために data-relpath が付与される。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        assert 'data-relpath="mismatch.txt"' in body
        assert 'data-relpath="ok.txt"' in body

    def test_detail_html_not_embedded_per_row(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """詳細 HTML を行ごとに埋め込まない (レポート肥大の原因だったため)。

        1,000 ファイル × 3 拠点で HTML の 87% がこの属性値だった。
        詳細は #report-data の JSON から JS が組み立てる。
        """
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        assert "data-detail-html" not in body

    def test_no_inline_detail_row_rendered(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """詳細はモーダル経由になったため、インラインの detail-row は出力されない。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        assert "detail-row" not in body

    def test_no_section_expand_controls(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """モーダル化により 全展開/全折りたたみ ボタンは削除済み。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        assert "data-action=" not in body
        assert "expand-all" not in body
        assert "collapse-all" not in body

    def test_report_data_has_one_entry_per_relpath(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """詳細データは相対パスをキーに 1 件ずつ。セクション間で重複しない。"""
        data = _extract_report_data(write_html(rich_ctx, tmp_path / "out.html"))
        # 和集合の 5 ファイル (ok, mismatch, missing_in_C, only_in_A, locked)
        assert set(data["rows"]) == {
            "ok.txt", "mismatch.txt", "missing_in_C.txt", "only_in_A.txt", "locked.txt",
        }
        assert data["rows"]["mismatch.txt"]["s"] == STATUS_HASH_MISMATCH

    def test_report_data_carries_normalized_roots(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """フルパスは JS が root + sep + relpath で組み立てるため、
        正規化済みルートとセパレータが拠点ごとに 1 回だけ入る。

        rich_ctx の拠点ルートは make_scan のデフォルト '/tmp/dummy' のため UNIX 形式。
        """
        data = _extract_report_data(write_html(rich_ctx, tmp_path / "out.html"))
        assert [loc["name"] for loc in data["locations"]] == ["拠点A", "拠点B", "拠点C"]
        for loc in data["locations"]:
            assert loc["root"] == "/tmp/dummy"
            assert loc["sep"] == "/"

    def test_report_data_records_entries_and_errors(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """各拠点のセルは [size, hash, mtime]、欠落は null、エラーは x に入る。"""
        data = _extract_report_data(write_html(rich_ctx, tmp_path / "out.html"))
        # only_in_A: 拠点A のみ存在
        only_a = data["rows"]["only_in_A.txt"]
        assert only_a["c"][0][1] == "hash_y"
        assert only_a["c"][1] is None and only_a["c"][2] is None
        assert "x" not in only_a  # エラーが無い行では省略
        # locked.txt: 拠点C が読み取り失敗
        locked = data["rows"]["locked.txt"]
        assert locked["c"] == [None, None, None]
        assert "PermissionError" in locked["x"][2]

    def test_detail_path_table_omits_folder_column(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """フォルダパス列は廃止 (ボタンには残るが列としては表示しない)。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        # 詳細テーブルのヘッダーは JS 側で組み立てる。列は 4 つで、フォルダパス列は無い。
        assert '"拠点", "状態", "ファイルパス", "操作"' in body
        assert '"フォルダパス"' not in body

    def test_modal_dialog_present(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """ページ末尾に <dialog id="detail-modal"> が1つ存在。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        assert '<dialog id="detail-modal"' in body
        assert 'id="detail-modal-title"' in body
        assert 'id="detail-modal-content"' in body
        # close ボタン
        assert 'class="modal-close"' in body

    def test_modal_script_present(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """showModal / close の呼び出しを含む JS が埋め込まれている。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        assert "showModal" in body
        assert "modal.close" in body

    def test_clipboard_script_with_fallback_present(
        self, rich_ctx: ReportContext, tmp_path: Path
    ):
        """新 API と旧 API フォールバック両方の呼び出しが script に含まれる。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        # 新 API
        assert "navigator.clipboard.writeText" in body
        # 旧 API フォールバック
        assert "document.execCommand" in body

    def test_size_column_highlighted_when_hash_skipped(self, tmp_path: Path):
        """hash_mode=smart でハッシュ未計算の不一致では、サイズ列の少数派を強調する。

        ハッシュが無いので通常のハッシュ強調が効かず、
        どの拠点が違うのかを示す手掛かりが消えてしまう。
        """
        a = make_scan("拠点A", files={"f.bin": make_entry(None, size=100)})
        b = make_scan("拠点B", files={"f.bin": make_entry(None, size=100)})
        c = make_scan("拠点C", files={"f.bin": make_entry(None, size=999)})
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b, c],
            comparison=compare([a, b, c]),
        )
        body = write_html(ctx, tmp_path / "out.html").read_text(encoding="utf-8")
        row = body.split('data-relpath="f.bin"')[1].split("</tr>")[0]
        # 少数派 (拠点C の 999) のサイズ列だけ強調され、多数派は素のまま
        assert "<td class='num cell-mismatch'>999</td>" in row
        assert row.count("cell-mismatch") == 1
        # ハッシュ欄は「未計算」表示
        assert row.count(SKIPPED_HASH_PLACEHOLDER) == 3

    def test_size_column_highlighted_in_excel_when_hash_skipped(self, tmp_path: Path):
        a = make_scan("拠点A", files={"f.bin": make_entry(None, size=100)})
        b = make_scan("拠点B", files={"f.bin": make_entry(None, size=100)})
        c = make_scan("拠点C", files={"f.bin": make_entry(None, size=999)})
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b, c],
            comparison=compare([a, b, c]),
        )
        out = write_excel(ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        ws = load_workbook(out)["全ファイル一覧"]
        # サイズ列は 6, 9, 12 (拠点ごとに3列)
        assert ws.cell(row=2, column=6).fill.fgColor.rgb in ("00000000", None)
        assert ws.cell(row=2, column=9).fill.fgColor.rgb in ("00000000", None)
        assert ws.cell(row=2, column=12).fill.fgColor.rgb == "00FFC7CE"
        assert ws.cell(row=2, column=13).value == SKIPPED_HASH_PLACEHOLDER

    def test_per_location_relpath_emitted_when_names_differ(self, tmp_path: Path):
        """拠点ごとに実ファイル名が違う場合だけ、行データに実パスを持たせる。

        macOS (NFD) と Windows (NFC) で同じファイルの名前の形が違うため、
        フルパスは拠点ごとの実際の名前から組み立てる必要がある。
        """
        nfc = unicodedata.normalize("NFC", "議事録_ガバナンス部会.docx")
        nfd = unicodedata.normalize("NFD", "議事録_ガバナンス部会.docx")
        entry = make_entry("h")
        a = make_scan("Win拠点", files={nfc: entry}, real_relpaths={nfc: nfc})
        b = make_scan("Mac拠点", files={nfc: entry}, real_relpaths={nfc: nfd})
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=compare([a, b]),
        )
        data = _extract_report_data(write_html(ctx, tmp_path / "out.html"))
        row = data["rows"][nfc]
        assert row["p"] == [nfc, nfd], "拠点ごとの実ファイル名が入っていない"
        assert unicodedata.is_normalized("NFD", row["p"][1])

    def test_per_location_relpath_omitted_when_names_match(self, tmp_path: Path):
        """名前が揃っていれば実パスは持たせない (行あたりのバイト数を増やさない)。"""
        a = make_scan("拠点A", files={"same.txt": make_entry("h")})
        b = make_scan("拠点B", files={"same.txt": make_entry("h")})
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=compare([a, b]),
        )
        data = _extract_report_data(write_html(ctx, tmp_path / "out.html"))
        assert "p" not in data["rows"]["same.txt"]

    def test_windows_roots_produce_backslash_paths(self, tmp_path: Path):
        """UNC / ドライブレターの拠点では、コピー用パスを Windows 形式で出す。

        本ツールの主対象は UNC 共有だが、パイプライン全体は POSIX でしか
        動かしていない。ルート正規化がレポートまで届いているかをここで見る。
        """
        entry = make_entry("h")
        a = make_scan("本社", files={"設計/仕様.xlsx": entry}, root=Path("//srv-hq/share/docs"))
        b = make_scan("支店", files={"設計/仕様.xlsx": entry}, root=Path("D:/shared/docs"))
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=compare([a, b]),
        )
        data = _extract_report_data(write_html(ctx, tmp_path / "out.html"))
        assert data["locations"] == [
            {"name": "本社", "root": "\\\\srv-hq\\share\\docs", "sep": "\\"},
            {"name": "支店", "root": "D:\\shared\\docs", "sep": "\\"},
        ]

    def test_report_is_deterministic(self, tmp_path: Path):
        """同じ入力からは同じレポートが出る (行順が実行ごとに揺れない)。

        フェーズ1を拠点ごとに並列化しているため、完了順が
        レポートの並びに漏れていないことを固定しておく。
        """
        ctx = _big_ctx(tmp_path, n_rows=200)
        first = write_html(ctx, tmp_path / "a.html").read_bytes()
        second = write_html(ctx, tmp_path / "b.html").read_bytes()
        assert first == second

        xa = write_excel(ctx, tmp_path / "a.xlsx")
        xb = write_excel(ctx, tmp_path / "b.xlsx")
        from openpyxl import load_workbook
        rows_a = list(load_workbook(xa)["全ファイル一覧"].iter_rows(values_only=True))
        rows_b = list(load_workbook(xb)["全ファイル一覧"].iter_rows(values_only=True))
        assert rows_a == rows_b

    def test_summary_paths_are_not_truncated_and_copyable(self, tmp_path: Path):
        """サマリーのパスは省略せず全体を出し、コピーできる。

        ファイル行と違ってモーダルが無いため、末尾を省略すると
        フルパスを確認する手段がなくなる。
        """
        long_root = "//server-with-a-very-long-name/share/部署/2026年度/資料一式"
        a = make_scan("拠点A", files={"x.txt": make_entry("h")})
        a.root = Path(long_root)
        b = make_scan("拠点B", files={"x.txt": make_entry("h")})
        config_path = tmp_path / "とても長いディレクトリ名" / "設定ファイル.yaml"
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=config_path,
            scans=[a, b],
            comparison=compare([a, b]),
        )
        out = write_html(ctx, tmp_path / "out.html")
        summary = out.read_text(encoding="utf-8").split('<section id="summary">')[1]
        summary = summary.split("</section>")[0]

        # 折り返し用クラスが付き、コピーボタンに全文が入る
        assert "wrap-path" in summary
        assert f'data-copy="{html_mod.escape(str(config_path), quote=True)}"' in summary
        assert f'data-copy="{html_mod.escape(long_root, quote=True)}"' in summary
        # 省略記号で切っていない
        assert str(config_path) in html_mod.unescape(summary)

    def test_error_and_dir_cells_are_not_truncated(self, tmp_path: Path):
        """エラー詳細とフォルダ構造差分にもモーダルが無いので省略しない。"""
        from scanner import ScanError

        long_msg = (
            "PermissionError: [Errno 13] Permission denied: "
            "'//server/share/部署/2026年度/とても長いファイル名の資料.xlsx'"
        )
        long_dir = "設計/2026年度/詳細設計/サブシステムA/インターフェース定義"
        a = make_scan(
            "拠点A",
            files={"x.txt": make_entry("h")},
            dirs={long_dir: object()},
            errors=[ScanError(relpath="部署/長いパス/資料.xlsx", message=long_msg)],
        )
        b = make_scan("拠点B", files={"x.txt": make_entry("h")})
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=compare([a, b]),
        )
        body = write_html(ctx, tmp_path / "out.html").read_text(encoding="utf-8")

        errors = body.split('<section id="errors">')[1].split("</section>")[0]
        assert errors.count("wrap-path") == 2  # 相対パスとメッセージ
        assert html_mod.escape(long_msg) in errors

        dirs = body.split('<section id="dirs">')[1].split("</section>")[0]
        assert "wrap-path" in dirs
        assert html_mod.escape(long_dir) in dirs

    def test_tables_have_filter_controls(self, rich_ctx: ReportContext, tmp_path: Path):
        """各ファイル表に絞り込み UI が付く。状態が1種類の表には状態選択を出さない。"""
        out = write_html(rich_ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")

        # 全ファイル一覧は複数の状態を含むので状態選択がある
        all_section = body.split('<section id="all">')[1].split("</section>")[0]
        assert 'class="filter-input"' in all_section
        assert 'class="filter-status"' in all_section
        for status in (STATUS_OK, STATUS_HASH_MISMATCH, STATUS_PARTIAL_MISSING):
            assert f'<option value="{status}">' in all_section

        # ハッシュ不一致セクションは状態が1種類なので選択は出さない
        mismatch = body.split('<section id="mismatch">')[1].split("</section>")[0]
        assert 'class="filter-input"' in mismatch
        assert 'class="filter-status"' not in mismatch

    def test_html_generation_does_not_hold_the_report_in_memory(self, tmp_path: Path):
        """HTML はファイルへ流しながら書く (行数に比例してメモリを使わない)。

        全体を 1 つの文字列に組み立てていた頃は 50,000 行 × 3 拠点で
        ピーク約 500MB だった。ここでは 2,000 行で、出力サイズより
        十分小さいピークに収まることを確認する。
        """
        ctx = _big_ctx(tmp_path, n_rows=2000)
        tracemalloc.start()
        try:
            out = write_html(ctx, tmp_path / "big.html")
            peak = tracemalloc.get_traced_memory()[1]
        finally:
            tracemalloc.stop()

        size = out.stat().st_size
        assert size > 1_000_000, "テストの前提が崩れている (出力が小さすぎる)"
        assert peak < size / 4, (
            f"ピーク {peak/1024**2:.1f}MB / 出力 {size/1024**2:.1f}MB — "
            "レポート全体をメモリに組み立てていないか確認"
        )

    def test_excel_generation_does_not_hold_the_sheet_in_memory(self, tmp_path: Path):
        """Excel は write_only で 1 行ずつ書く (シート全体を保持しない)。

        通常モードは 50,000 行 × 3 拠点で約 236MB を使っていた。
        """
        ctx = _big_ctx(tmp_path, n_rows=2000)
        tracemalloc.start()
        try:
            write_excel(ctx, tmp_path / "big.xlsx")
            peak = tracemalloc.get_traced_memory()[1]
        finally:
            tracemalloc.stop()

        # write_only なら行数によらずほぼ一定。通常モードだと 2,000 行で 10MB 前後になる
        assert peak < 5 * 1024 * 1024, (
            f"ピーク {peak/1024**2:.1f}MB — write_only が外れていないか確認"
        )

    def test_report_size_scales_modestly_with_row_count(self, tmp_path: Path):
        """1 行あたりの出力バイト数に上限を設ける (肥大の再発防止)。

        詳細 HTML を行ごとに埋め込んでいた頃は 3 拠点で 1 行 5.3KB あった。
        JSON 集約後はテーブルセルが支配的になり、1 行 2KB を大きく下回る。
        """
        n = 300
        files = {f"dir{i // 20}/file_{i:04d}.txt": make_entry(f"h{i:060d}") for i in range(n)}
        a = make_scan("拠点A", files=files)
        b = make_scan("拠点B", files=files)
        c = make_scan("拠点C", files=files)
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b, c],
            comparison=compare([a, b, c]),
        )
        out = write_html(ctx, tmp_path / "big.html")
        per_row = out.stat().st_size / n
        assert per_row < 2048, f"1 行あたり {per_row:.0f}B — 詳細データが重複していないか確認"

    @pytest.mark.parametrize(
        "relpath, expected",
        [
            # XML に書けない制御文字。openpyxl が例外を投げてレポートが全滅していた
            ("資料\x07ベル.docx", "資料\\x07ベル.docx"),
            # 不正な UTF-8 のファイル名 (surrogateescape)。
            # HTML の書き出しが UnicodeEncodeError、Excel は壊れた xlsx を吐いていた
            ("資料\udcff壊れ.docx", "資料\\udcff壊れ.docx"),
            ("報告\x01\udcfe.docx", "報告\\x01\\udcfe.docx"),
        ],
    )
    def test_unrepresentable_filenames_do_not_break_the_report(
        self, tmp_path: Path, relpath: str, expected: str
    ):
        """出力できない文字を含む名前でも、レポートを落とさず見える表記で出す。

        共有フォルダにはレガシー機器が付けた名前や文字コード不一致のファイルが
        実在する。1 ファイルのためにレポート全体を失うほうが困る。
        """
        entry = make_entry("a" * 64)
        a = make_scan("拠点A", files={relpath: entry})
        b = make_scan("拠点B", files={relpath: entry})
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=compare([a, b]),
        )

        # HTML: 生成でき、表の照合キーと詳細データのキーが一致する
        out = write_html(ctx, tmp_path / "out.html")
        body = out.read_text(encoding="utf-8")
        data = _extract_report_data(out)
        assert list(data["rows"]) == [expected]
        attr = re.search(r'data-relpath="([^"]*)"', body).group(1)
        assert html_mod.unescape(attr) == expected, "表とJSONでキーが食い違うと詳細が引けない"

        # Excel: 生成でき、読み戻せる (壊れた xlsx を吐かない)
        xlsx = write_excel(ctx, tmp_path / "out.xlsx")
        from openpyxl import load_workbook
        assert load_workbook(xlsx)["全ファイル一覧"].cell(row=2, column=2).value == expected

    def test_html_escapes_special_chars_in_paths(self, tmp_path: Path):
        """パス名に <script> を混ぜても素通りしないこと (HTMLエスケープ)。"""
        a = make_scan("A", files={"<script>.txt": make_entry("h")})
        b = make_scan("B", files={"<script>.txt": make_entry("h")})
        comparison = compare([a, b])
        ctx = ReportContext(
            started_at=datetime(2026, 1, 1),
            finished_at=datetime(2026, 1, 1),
            config_path=tmp_path / "c.yaml",
            scans=[a, b],
            comparison=comparison,
        )
        out = write_html(ctx, tmp_path / "evil.html")
        body = out.read_text(encoding="utf-8")
        # ファイル名由来の <script>.txt は素のまま出ない (エスケープされている)
        assert "<script>.txt" not in body
        assert "&lt;script&gt;.txt" in body
